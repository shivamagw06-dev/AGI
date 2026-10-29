"""Read the supplied research workbook; never modify it or invent missing inputs."""
import argparse, collections, datetime, hashlib, json, pathlib, re
import openpyxl

parser = argparse.ArgumentParser()
parser.add_argument('workbook')
parser.add_argument('--output', default='server/data/private-capital-watchlist.json')
args = parser.parse_args()
book = openpyxl.load_workbook(args.workbook, data_only=True)
def value(v):
    if isinstance(v, (datetime.datetime, datetime.date)): return v.strftime('%Y-%m-%d')
    if isinstance(v, str) and (v.lower().strip() in ['n/a', 'n.a.', 'n.m.', '-'] or v.lower().startswith('n/a (')): return None
    return v
def records(sheet, header, start, end, keys):
    return [dict(zip(keys, [value(c.value) for c in row])) | {'source': {'sheet': sheet, 'row': row[0].row}}
            for row in book[sheet].iter_rows(min_row=start, max_row=end, max_col=len(keys)) if row[0].value is not None]
companies = records('Master_100',1,2,101,['number','name','legalEntity','sector','subsector','hq','category','signal','dealUsdM','valuationUsdM','valuationDate','raisedUsdM','rounds','firstFundingDate','fundingDate','fundingAgeMonths','revenueInrCr','fiscalYear','priorRevenueInrCr','priorFiscalYear','growth','revenueUsdM','raisedRevenue','status','provider','fundingFlag'])
gaps = collections.defaultdict(list)
for name, issue, fix in book['Data_Gaps'].iter_rows(min_row=2,values_only=True):
    if name: gaps[name].append({'issue':issue,'nextStep':fix})
for c in companies:
    c['id'] = re.sub(r'[^a-z0-9]+','-',c['name'].lower()).strip('-')
    c['gaps'] = gaps[c['name']]
    c['review'] = 'Source review pending'
    c['notes'] = []
    if c['dealUsdM'] is not None and c['raisedUsdM'] is not None and c['dealUsdM'] > c['raisedUsdM']:
        c['notes'].append('Deal size exceeds recorded cumulative funding. Reconcile equity, debt, secondary consideration and incomplete history before comparison.')
    if c['fundingDate'] and c['valuationDate'] and c['valuationDate'] > c['fundingDate']:
        c['notes'].append('Valuation date is newer than the recorded funding date. The funding timeline needs reconciliation.')
    if c['valuationDate'] and c['valuationDate'] < '2024-10-29': c['notes'].append('Reported valuation is more than two years old; it is not a current transaction price.')
    if c['fiscalYear'] in ['FY2021','FY2022']: c['notes'].append('Financials are historical. More recent annual accounts are required.')
    if c['sector']=='AI / Software' and c['subsector']=='Space launch': c['sector']='Space / Deep tech'
    if c['subsector']=='Fantasy sports / gaming': c['sector']='Gaming'
transactions = records('Precedent_Txns',3,4,24,['date','company','business','buyer','sellers','stake','considerationNativeM','currency','considerationInrBn','equityValueInrBn','reportedEquityValueInrBn','revenueInrBn','revenuePeriod','valueRevenue','netIncomeInrBn','pe','payment','provider'])
comparables = records('Listed_Comps',6,7,24,['company','group','marketCapInrBn','netCashInrBn','evInrBn','revenueInrBn','growth','cagr','ebitdaInrBn','ebitdaMargin','fy26EbitdaMargin','netMargin','evRevenue','evEbitda','pe','netCashShare','stakesInrBn','coreEvInrBn','coreEvEbitda'])
for x in comparables:
    x['period'] = value(book['Listed_Financials'].cell(x['source']['row']-5,15).value)
    if x['company'].startswith('Info Edge'):
        x['coreEvInrBn']=None; x['coreEvEbitda']=None
secondary = records('Secondary_Map',7,8,19,['company','fundingDate','fundingAgeMonths','valuationUsdM','raisedUsdM','firstInvestorYear','holdingYears','historicalInvestors','crossoverFlag','exitStatus','score','rank'])
for x in secondary:
    x.pop('score'); x.pop('rank')
    if x['valuationUsdM']==0: x['valuationUsdM']=None
management = records('Mgmt_Scorecard',9,10,19,['company','ceo','founderCeo','cfo','cfoPresent','irPresent','formerExecutives','ceoChange','score','observations'])
for x in management:
    x.pop('score')
    if x['cfo'] and ('Not listed' in x['cfo'] or x['cfo'].startswith('n/a')): x['cfoPresent']=None
checks = [
 {'company':'Sarvam AI','date':'2026-06-15','title':'Series B first close','text':'Company announcement reports $234m first close of a $300m Series B at a $1.5bn post-money valuation. The workbook funding history has not been reconciled to this event.','url':'https://www.sarvam.ai/announcing-series-b'},
 {'company':'Neysa Networks','date':'2026-02-16','title':'Equity and debt financing','text':'Blackstone describes up to $600m equity and an intended additional $600m debt financing, subject to documentation. The $1.2bn package must not be treated as completed equity or buyout consideration.','url':'https://www.blackstone.com/news/press/blackstone-leads-funding-of-over-1-billion-to-neysa-to-work-towards-building-indias-leading-ai-infrastructure-platform/'},
 {'company':'Nxtra Data','date':'2026-03-30','title':'Announced growth investment','text':'Airtel announces a $1bn investment and approximately $3.1bn post-closing valuation, subject to approvals. This is announcement evidence; completed funding remains to be reconciled.','url':'https://www.airtel.in/press-release/03-2026/airtel-announces-us1-billion-investment-in-nxtra-led-by-alpha-wave-global-and-existing-investor-carlyle-bharti-airtel-will-also-participate/'}]
data={'asOf':'2026-10-29','workbook':'India_PE_VC_Top100_Watchlist.xlsx','sha256':hashlib.sha256(pathlib.Path(args.workbook).read_bytes()).hexdigest(),'fxInrPerUsd':88,'companies':companies,'transactions':transactions,'comparables':comparables,'secondary':secondary,'management':management,'sourceChecks':checks,'methodology':['Curated company watchlist, not a ranked Top 100 or investment recommendation.','Workbook figures are reported research inputs unless a linked source check explicitly states otherwise.','Round valuations are historical observations, not executable prices or fair values.','Funding currency and transaction scope require reconciliation. Cumulative funding is not aggregated here.','Secondary ownership and leadership observations require current verification. Numerical heuristic scores are withheld.','Static research snapshot. It does not refresh from PitchBook or Capital IQ.']}
out=pathlib.Path(args.output); out.parent.mkdir(parents=True,exist_ok=True); out.write_text(json.dumps(data,indent=2,ensure_ascii=False)+'\n')
print(f'Imported {len(companies)} companies, {len(transactions)} transactions, {len(comparables)} listed peers.')
