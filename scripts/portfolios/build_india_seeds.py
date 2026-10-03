import openpyxl,json,math,sys
from pathlib import Path
from collections import defaultdict
w=openpyxl.load_workbook(sys.argv[1],data_only=True)
out=[]
for s in list(w)[1:]:
 rows=[r for r in s.iter_rows(min_row=7,max_col=12,values_only=True) if r[2]]
 hs=[dict(name=r[1],symbol=r[2].split(':')[-1],sector=r[3],sourceRank=r[0],score=r[4],marketCapCr=r[11]) for r in rows]
 capacity=defaultdict(float)
 for h in hs:capacity[h['sector']]+=4 if h['marketCapCr']<10000 else 7
 target=min(100,sum(min(25,v) for v in capacity.values()))
 base=[5-.5*(i//5) for i in range(len(hs))];weights=[0.0]*len(hs)
 for step in range(2000):
  sector=defaultdict(float)
  for h,v in zip(hs,weights):sector[h['sector']]+=v
  remaining=target-sum(weights)
  if remaining<1e-9:break
  available=[i for i,h in enumerate(hs) if weights[i]<(4 if h['marketCapCr']<10000 else 7)-1e-9 and sector[h['sector']]<25-1e-9]
  if not available:raise Exception('Infeasible caps')
  den=sum(base[i] for i in available);proposal={i:remaining*base[i]/den for i in available}
  scale=1.0
  for i in available:scale=min(scale,((4 if hs[i]['marketCapCr']<10000 else 7)-weights[i])/proposal[i])
  for sec in sector:
   inc=sum(proposal[i] for i in available if hs[i]['sector']==sec)
   if inc:scale=min(scale,(25-sector[sec])/inc)
  for i in available:weights[i]+=proposal[i]*scale
 else:raise Exception('Failed convergence')
 bps=[math.floor(x*100+1e-7) for x in weights]
 while sum(bps)<round(target*100):
  sector=defaultdict(int)
  for h,v in zip(hs,bps):sector[h['sector']]+=v
  valid=[i for i,h in enumerate(hs) if bps[i]<(400 if h['marketCapCr']<10000 else 700) and sector[h['sector']]<2500]
  i=max(valid,key=lambda i:weights[i]*100-bps[i]);bps[i]+=1
 for h,b in zip(hs,bps):h['weight']=b/100
 if target<100:hs.append(dict(name='Nippon India ETF Gold BeES',symbol='GOLDBEES',sector='Gold',assetType='gold_etf',weight=100-target))
 out.append(dict(id='in-'+s.title.lower(),name='AGI '+s.title,market='india',category='Factor strategies',description=s.cell(3,1).value+'. Research allocation from the filtered shortlist; factor scores and analyst coverage are not fully verified.',asOf='2026-09-30',holdings=hs,cashWeight=0,incomplete=False,customized=True,trackingStart='2026-10-05',weightMethod='Rank tiers by retained order: 5/4.5/4/3.5/3 points per five stocks, proportionally allocated subject to 7% stock, 4% below INR 10,000cr, and 25% source-sector caps. Rounded to 0.01% within caps.',sourceName='AGI filtered shortlist; supplied sector and market-cap snapshot',sourceAsOf='2026-09-30'))
open(Path(__file__).resolve().parents[2]/'server/data/indiaPortfolioSeeds.json','w').write(json.dumps(out,indent=2)+'\n')
for p in out:print(p['name'],len(p['holdings']),sum(h['weight'] for h in p['holdings']))
