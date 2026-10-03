"""Extract India long/short deciles. Arguments: JPM workbook, ticker workbook, Upstox NSE master gzip."""
import sys,json,gzip
from pathlib import Path
import openpyxl
source,mapping,master=sys.argv[1:]
w=openpyxl.load_workbook(source,data_only=True)
m=openpyxl.load_workbook(mapping,data_only=True).active
symbols={r[0]:str(r[13]).replace('NSEI:','') for r in list(m.values)[1:]}
symbols['OINL IN']=symbols['OILN IN'];symbols['VAML IN']='VAML'
instruments=json.load(gzip.open(master)); factors=['Composite','Value','Growth','Quality','Momentum','Low Vol']
out=[]
for side,col in [('long',2),('short',9)]:
 categories=[]; unique={}
 for factor in factors:
  s=w[factor];holdings=[]
  for r in range(5,s.max_row+1):
   ticker,name,market,sector,score=[s.cell(r,col+i).value for i in range(5)]
   if market!='India':continue
   symbol=symbols[ticker];hits=[x for x in instruments if x.get('trading_symbol')==symbol and x.get('instrument_type')=='EQ']
   assert len(hits)==1,(ticker,symbol,len(hits))
   h={'symbol':symbol,'sourceTicker':ticker,'name':hits[0]['name'],'sector':sector,'instrumentKey':hits[0]['instrument_key'],'score':score,'sourceCell':f'{factor}!{s.cell(r,col+4).coordinate}'}
   holdings.append(h)
  for h in holdings:
   h['weight']=100/len(holdings)
   u=unique.setdefault(h['symbol'],{k:v for k,v in h.items() if k not in ['score','sourceCell','weight']})
   u['weight']=u.get('weight',0)+h['weight']/6
   u.setdefault('factors',[]).append(factor)
  categories.append({'name':factor,'weight':100/6,'definition':s['O1'].value,'holdings':holdings})
 p={'id':f'in-conviction-{side}','name':'Conviction Long' if side=='long' else 'Conviction Short — Research','market':'india','category':'Conviction','conviction':True,'direction':side,'categories':categories,'holdings':list(unique.values()),'cashWeight':0,'incomplete':False,'asOf':'2026-10-03','trackingStart':'2026-10-05','sourceName':Path(source).name,'sourceAsOf':'2026-09-01','description':f'India {side}-decile classifications across six factor categories. Equal category allocations, equal stock weights within each category; shared stocks remain in every applicable category.','weightMethod':'AGI allocation: one-sixth to each category, then equal weights within that category. Shared-stock weights are added for the total. Fixed launch units; no automatic rebalancing. Composite classification comes from JPM, not an average of visible scores.','researchOnly':True}
 assert abs(sum(h['weight'] for h in p['holdings'])-100)<1e-9
 out.append(p)
Path('server/data/indiaConvictionPortfolios.json').write_text(json.dumps(out,indent=2)+'\n')
print([(p['id'],len(p['holdings']),sum(len(c['holdings']) for c in p['categories'])) for p in out])
